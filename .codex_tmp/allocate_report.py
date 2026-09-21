from pathlib import Path
from zipfile import ZipFile
from copy import deepcopy
from lxml import etree

source=Path('D:/Hk1 năm 3/LT_KDCLPM/TD01_NhomXX.docx')
target=Path('C:/Users/vohon/petnova_app/docs/TD01_Nhom09.docx')
ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
W='{'+ns['w']+'}'
with ZipFile(source) as zin:
    xml=etree.fromstring(zin.read('word/document.xml'))
    table=xml.find('.//w:tbl',ns)
    rows=table.findall('w:tr',ns)
    assert len(rows)==5
    def put_paragraph(p,text):
        oldr=p.find('w:r',ns)
        rpr=deepcopy(oldr.find('w:rPr',ns)) if oldr is not None and oldr.find('w:rPr',ns) is not None else None
        for child in list(p):
            if child.tag != W+'pPr': p.remove(child)
        r=etree.SubElement(p,W+'r')
        if rpr is not None: r.append(rpr)
        t=etree.SubElement(r,W+'t'); t.text=text
    def put_cell(cell,text):
        p=cell.find('w:p',ns)
        for child in list(cell):
            if child.tag!=W+'tcPr' and child is not p: cell.remove(child)
        put_paragraph(p,text)
    head=rows[0].findall('w:tc',ns)
    put_cell(head[1],'Phần việc phụ trách đề xuất')
    put_cell(head[2],'Tỷ lệ đóng góp dự kiến')
    tasks=[
        'Đăng ký, đăng nhập và hồ sơ khách hàng; quản lý thú cưng và tải ảnh. Kiểm thử, viết báo cáo phần phụ trách.',
        'Đặt lịch, cập nhật trạng thái dịch vụ; giao diện nhân viên; thanh toán tiền mặt và PayOS. Kiểm thử, viết báo cáo phần phụ trách.',
        'Giao diện bác sĩ; bệnh án, điều trị, tiêm chủng và thông báo. Kiểm thử, viết báo cáo phần phụ trách.',
        'Giao diện quản trị; quản lý tài khoản, nhân sự, gói dịch vụ và thống kê. Phối hợp tích hợp, tổng hợp kiểm thử và báo cáo.'
    ]
    for row,text in zip(rows[1:],tasks):
        cells=row.findall('w:tc',ns)
        put_cell(cells[1],text)
        put_cell(cells[2],'25%')
    changed=0
    for p in xml.findall('.//w:body/w:p',ns):
        text=''.join(p.itertext()) if False else ''.join(p.xpath('.//w:t/text()',namespaces=ns))
        if text=='Thành viên và đánh giá đóng góp':
            put_paragraph(p,'Phân công và đóng góp dự kiến'); changed+=1
        elif text.startswith('Đánh giá theo khối lượng công việc, chất lượng sản phẩm'):
            put_paragraph(p,'Phân công và tỷ lệ trên là đề xuất, tổng cộng 100%. Mỗi thành viên kiểm thử và viết báo cáo cho phần phụ trách. Nhóm xác nhận mức đóng góp thực tế theo khối lượng, chất lượng, tiến độ và phối hợp trước khi nộp.'); changed+=1
    assert changed==2
    updated=etree.tostring(xml,xml_declaration=True,encoding='UTF-8',standalone=True)
    with ZipFile(target,'w') as zout:
        for item in zin.infolist():
            zout.writestr(item,updated if item.filename=='word/document.xml' else zin.read(item.filename))
with ZipFile(source) as old, ZipFile(target) as new:
    assert new.testzip() is None
    assert old.namelist()==new.namelist()
    assert all(old.read(n)==new.read(n) for n in old.namelist() if n!='word/document.xml')
    print('Verified: all package parts except document text preserved.')
from docx import Document
d=Document(target)
assert len(d.tables)==4
for row in d.tables[0].rows: print(' | '.join(c.text for c in row.cells))
print(target)
