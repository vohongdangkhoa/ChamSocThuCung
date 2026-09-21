from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from pathlib import Path

root=Path('C:/Users/vohon/petnova_app')
doc=Document()
sec=doc.sections[0]
sec.page_width=Cm(21); sec.page_height=Cm(29.7)
sec.top_margin=Cm(2); sec.bottom_margin=Cm(2)
sec.left_margin=Cm(2.3); sec.right_margin=Cm(2)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
    s=doc.styles[name]; s.font.name='Times New Roman'; s.font.color.rgb=RGBColor(0,0,0)
    s.font.size=Pt(12)
    s.paragraph_format.space_after=Pt(6)
    s.paragraph_format.line_spacing=1.12
doc.styles['Title'].font.size=Pt(23)
doc.styles['Title'].font.bold=True
doc.styles['Heading 1'].font.size=Pt(16)
doc.styles['Heading 2'].font.size=Pt(13)
for n in ['Heading 1','Heading 2']:
    doc.styles[n].paragraph_format.space_before=Pt(10)
    doc.styles[n].paragraph_format.keep_with_next=True
def p(t,style=None): return doc.add_paragraph(t,style)
def h(t): doc.add_heading(t,1)
def sub(t): doc.add_heading(t,2)
def page(): doc.add_page_break()
def table(headers, rows, widths):
    t=doc.add_table(rows=1, cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
    for i,w in enumerate(widths): t.columns[i].width=Cm(w)
    for i,s in enumerate(headers): t.rows[0].cells[i].text=s
    for row in rows:
        for c,s in zip(t.add_row().cells,row): c.text=s
    for ri,row in enumerate(t.rows):
        pr=row._tr.get_or_add_trPr(); pr.append(OxmlElement('w:cantSplit'))
        if ri==0: pr.append(OxmlElement('w:tblHeader'))
        for i,c in enumerate(row.cells):
            c.width=Cm(widths[i]); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cp=c._tc.get_or_add_tcPr(); borders=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                e=OxmlElement('w:'+edge); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'4'); e.set(qn('w:color'),'D9D9D9'); borders.append(e)
            cp.append(borders)
            shade=OxmlElement('w:shd'); shade.set(qn('w:fill'),'DCE6EF' if ri==0 else 'FFFFFF'); cp.append(shade)
            mar=OxmlElement('w:tcMar')
            for edge in ['top','left','bottom','right']:
                e=OxmlElement('w:'+edge); e.set(qn('w:w'),'85'); e.set(qn('w:type'),'dxa'); mar.append(e)
            cp.append(mar)
            for para in c.paragraphs:
                para.paragraph_format.space_after=Pt(2); para.paragraph_format.line_spacing=1.05
                for r in para.runs: r.font.size=Pt(11); r.bold=ri==0
    p('')

p('BÁO CÁO TIẾN ĐỘ WEBSITE PETNOVA','Title')
p('Quản lý chăm sóc và sức khỏe thú cưng','Subtitle')
p('Đợt báo cáo  TD01     |     Ngày 21 tháng 09 năm 2026')
p('Nhóm: ...............     Lớp: ........................................................')
p('Học phần: ................................................................................')
p('Giảng viên: ..............................................................................')
sub('Tổng quan tiến độ')
p('PetNoVa đã có giao diện web cho bốn vai trò, kết nối API và các luồng quản lý thú cưng, đặt lịch, thanh toán, bệnh án và tiêm chủng. Bản dựng giao diện đã thực hiện thành công. Giai đoạn tiếp theo tập trung kiểm thử tích hợp, hoàn thiện kiểm soát quyền truy cập và tính nhất quán dữ liệu trước khi nghiệm thu.')
p('Mức đánh giá trong báo cáo: “Đã xây dựng” là đã có giao diện hoặc mã xử lý tương ứng; không đồng nghĩa với đã kiểm thử đạt toàn bộ nghiệp vụ. Các tích hợp Firebase, PayOS, Cloudinary và email cần được kiểm chứng trong môi trường chạy thực tế.')
sub('Thành viên và đánh giá đóng góp')
table(['Họ tên và MSSV','Phần việc thực hiện','Đóng góp và đánh giá'],[
('................................\n................................','................................\n................................','.......... %\n................................'),
('................................\n................................','................................\n................................','.......... %\n................................'),
('................................\n................................','................................\n................................','.......... %\n................................'),
('................................\n................................','................................\n................................','.......... %\n................................')],[5.3,6.1,5.3])
p('Đánh giá theo khối lượng công việc, chất lượng sản phẩm, tiến độ và phối hợp. Thống nhất tỷ lệ đóng góp giữa các thành viên, tổng cộng 100%.')

page(); h('A Mục tiêu phần mềm')
p('PetNoVa hỗ trợ cơ sở chăm sóc thú cưng tổ chức lịch phục vụ và lưu trữ thông tin chăm sóc tập trung. Khách hàng có thể quản lý hồ sơ thú cưng, đặt dịch vụ, theo dõi thanh toán và xem lịch sử sức khỏe trên website.')
sub('Mục tiêu cụ thể')
p('1. Số hóa hồ sơ chủ nuôi và thú cưng, giúp tra cứu thông tin liên hệ, đặc điểm và lịch sử chăm sóc.')
p('2. Quản lý lịch hẹn xuyên suốt từ đặt lịch, xác nhận, thực hiện dịch vụ đến hoàn thành hoặc hủy.')
p('3. Ghi nhận thanh toán tiền mặt và tích hợp chuyển khoản PayOS, đồng thời theo dõi trạng thái giao dịch.')
p('4. Lưu bệnh án, điều trị, mũi tiêm và ngày nhắc lại để khách hàng và bác sĩ theo dõi sức khỏe thú cưng.')
p('5. Cung cấp công cụ quản trị tài khoản, nhân sự, gói dịch vụ và số liệu tổng quan cho cơ sở.')
sub('Phạm vi và nền tảng')
p('Phiên bản web sử dụng React, JavaScript và CSS với Vite. Máy chủ sử dụng ASP.NET Core trên .NET 10; dữ liệu được quản lý bằng SQL Server thông qua Entity Framework Core. Firebase phục vụ xác thực, Cloudinary lưu ảnh, PayOS xử lý thanh toán chuyển khoản và SMTP gửi email khôi phục mật khẩu.')
h('B Các vai trò trong hệ thống')
table(['Vai trò','Chức năng theo giao diện hiện có'],[
('Khách hàng\nCUSTOMER','Đăng ký, đăng nhập; quản lý hồ sơ và thú cưng; đặt hoặc hủy lịch; xem thanh toán, bệnh án, tiêm chủng và thông báo.'),
('Nhân viên\nSTAFF','Tìm kiếm, lọc lịch hẹn; xác nhận và cập nhật tiến độ dịch vụ; xác nhận tiền mặt; kiểm tra giao dịch PayOS.'),
('Bác sĩ thú y\nVET','Tra cứu thú cưng và lịch sử sức khỏe; lập bệnh án; ghi nhận vaccine, ngày tiêm và ngày nhắc lại.'),
('Quản trị viên\nADMIN','Xem tổng quan doanh thu và lịch; quản lý tài khoản, vai trò, trạng thái, nhân sự và gói dịch vụ.')],[4,12.7])
p('Quyền nghiệp vụ cần được kiểm soát tại API. Việc chia màn hình theo vai trò trên website chỉ là một lớp điều hướng và chưa thay thế kiểm tra quyền phía máy chủ.')

page(); h('C Các chức năng cơ bản và chi tiết')
table(['Nhóm chức năng','Nội dung đã xây dựng','Mức hiện tại'],[
('Tài khoản','Đăng ký và đăng nhập Firebase; tải hồ sơ SQL; chặn hồ sơ không ACTIVE; chuyển giao diện theo vai trò.','Có giao diện và xử lý'),
('Khôi phục mật khẩu','Nhập số điện thoại; gửi OTP tới email liên kết; xác minh OTP; đặt mật khẩu mới bằng vé dùng một lần.','Có xử lý; cần thử email thực tế'),
('Hồ sơ thú cưng','Thêm, sửa tên, loài, giống, giới tính, ngày sinh, cân nặng, tình trạng sức khỏe; tải ảnh thú cưng.','Có giao diện và API'),
('Đặt lịch dịch vụ','Chọn thú cưng, gói đang hoạt động, ngày, giờ, hình thức trả tiền và ghi chú; xem chi tiết; hủy lịch hợp lệ.','Có luồng; cần hoàn thiện giao dịch dữ liệu'),
('Vận hành lịch','Tìm và lọc lịch; chuyển trạng thái theo thứ tự; chặn hoàn thành nếu chưa ghi nhận thanh toán.','Có giao diện và API'),
('Thanh toán','Tạo khoản thanh toán; xác nhận tiền mặt; tạo link, đồng bộ và nhận webhook PayOS.','Có tích hợp; cần kiểm thử đầu cuối'),
('Hồ sơ sức khỏe','Bác sĩ thêm chẩn đoán, điều trị, ghi chú; ghi nhận tiêm chủng; khách xem lịch sử theo thú cưng.','Có giao diện và API'),
('Thông báo','Xem thông báo theo tài khoản; đánh dấu một hoặc tất cả đã đọc; mở lịch hẹn liên quan.','Có giao diện và API'),
('Quản trị','Thống kê khoản PAID; quản lý vai trò và trạng thái tài khoản; thêm, sửa nhân sự và gói dịch vụ.','Có giao diện và API')],[3.2,9.5,4])
sub('Quan hệ dữ liệu chính')
p('Một tài khoản khách hàng gắn với nhiều thú cưng và lịch hẹn. Lịch hẹn liên kết thú cưng, chi tiết dịch vụ và khoản thanh toán. Bệnh án và tiêm chủng liên kết thú cưng với nhân sự phụ trách. Thông báo gắn với người nhận và có thể dẫn tới lịch hẹn liên quan.')
p('Tổng quan doanh thu hiện cộng các khoản có trạng thái PAID. Cần đối soát thêm với quy trình hoàn tiền trước khi dùng số liệu cho báo cáo vận hành chính thức.')

page(); h('D Các ràng buộc trong xử lý')
sub('Các quy tắc đã thể hiện trong mã')
p('Tài khoản: giao diện chỉ mở cổng chức năng khi hồ sơ ở trạng thái ACTIVE và vai trò thuộc CUSTOMER, STAFF, VET hoặc ADMIN. Hồ sơ và vai trò được tải từ máy chủ sau đăng nhập Firebase.')
p('Đặt lịch: biểu mẫu yêu cầu thú cưng, gói dịch vụ, ngày và giờ; không cho chọn ngày trước hôm nay. API tạo lịch ở trạng thái PENDING, tự đặt thời gian tạo và bỏ thông tin phân công nhân viên do phía khách gửi lên.')
p('Trạng thái: luồng chính là PENDING → CONFIRMED → IN_PROGRESS → COMPLETED. Chỉ hủy lịch khi PENDING hoặc CONFIRMED. Lịch COMPLETED hoặc CANCELLED không được chuyển tiếp. Hoàn thành dịch vụ yêu cầu có khoản thanh toán PAID.')
p('Thanh toán: số tiền phải lớn hơn 0 và lịch hẹn phải tồn tại. Khoản mới được đặt PENDING. Xác nhận thủ công chỉ áp dụng CASH; tạo link PayOS chỉ dành cho BANK_TRANSFER. Webhook có xử lý kiểm tra dữ liệu giao dịch trước khi ghi nhận trả tiền.')
p('Ảnh: máy chủ chấp nhận JPEG, PNG hoặc WebP, kiểm tra nội dung file và giới hạn 5 MB. Một số biểu mẫu đang cho phép tới 6 MB, cần thống nhất để tránh khách chọn ảnh rồi bị API từ chối.')
p('Khôi phục mật khẩu: áp dụng cho khách hàng ACTIVE; chuẩn hóa số điện thoại; OTP gửi đến email liên kết. Luồng có thời hạn, kiểm soát thử sai và vé đặt lại dùng một lần. API giới hạn 20 yêu cầu trong 5 phút theo IP; phía web yêu cầu HTTPS ngoài môi trường loopback.')
sub('Ràng buộc cần hoàn thiện trước nghiệm thu')
p('Phân quyền và sở hữu dữ liệu: bổ sung kiểm tra danh tính, vai trò và quyền trên từng bản ghi cho các API thú cưng, lịch hẹn, thanh toán, bệnh án và tiêm chủng. Các controller này chưa thể hiện cơ chế kiểm tra tương đương nhóm tài khoản và tải ảnh.')
p('Tính nhất quán: hiện website tạo lịch, chi tiết dịch vụ và thanh toán bằng ba yêu cầu riêng. Cần gộp thành một giao dịch phía máy chủ hoặc có cơ chế khôi phục khi bước giữa thất bại; giá tiền cần được tính lại từ danh mục dịch vụ tại máy chủ.')
p('Đồng thời và lịch phục vụ: thay cách sinh mã dựa trên số bản ghi cộng một bằng cơ chế an toàn khi nhiều người thao tác. Bổ sung quy tắc chống trùng lịch, kiểm tra năng lực phục vụ và kiểm tra ngày giờ tại API.')
p('Triển khai: giới hạn CORS theo tên miền chính thức, cấu hình HTTPS, bảo vệ thông tin kết nối và kiểm chứng các tích hợp bên ngoài trước khi công bố website.')

page(); h('E Quy trình xử lý chung và riêng')
sub('Quy trình chung của hệ thống')
p('Người dùng đăng nhập Firebase → website tải hồ sơ từ API → kiểm tra trạng thái và vai trò → mở cổng chức năng tương ứng → người dùng thực hiện nghiệp vụ → API xử lý và lưu SQL Server → website tải lại dữ liệu, hiển thị kết quả hoặc lỗi.')
sub('Quy trình đặt lịch và phục vụ')
p('Bước 1. Khách hàng tạo hoặc chọn thú cưng, chọn gói dịch vụ đang hoạt động, ngày giờ và phương thức thanh toán.')
p('Bước 2. Website kiểm tra dữ liệu nhập rồi lần lượt tạo lịch PENDING, chi tiết dịch vụ và khoản thanh toán PENDING. Nếu bước sau thất bại, cần xử lý phần dữ liệu đã tạo trước khi thử lại.')
p('Bước 3. Nhân viên xem danh sách, xác nhận lịch thành CONFIRMED và chuyển sang IN_PROGRESS khi bắt đầu phục vụ.')
p('Bước 4. Ghi nhận thanh toán. Khi đã có khoản PAID, nhân viên mới chuyển lịch sang COMPLETED. Các thay đổi trạng thái tạo thông báo cho khách hàng.')
p('Nhánh hủy: khách có thể hủy lịch ở trạng thái PENDING hoặc CONFIRMED. Hủy lịch và hoàn tiền là hai nghiệp vụ cần được đối chiếu riêng; hủy lịch không tự chứng minh đã hoàn tiền cho khách.')
sub('Quy trình thanh toán riêng')
p('Tiền mặt: tạo khoản CASH/PENDING → khách trả tại quầy → nhân viên xác nhận → khoản chuyển PAID và ghi ngày thanh toán.')
p('Chuyển khoản: tạo khoản BANK_TRANSFER/PENDING → yêu cầu link PayOS → khách mở link và thanh toán → máy chủ nhận webhook hoặc truy vấn trạng thái PayOS → cập nhật kết quả và thông báo. Không lấy việc khách quay lại trang làm bằng chứng duy nhất đã trả tiền.')
sub('Quy trình khám và tiêm chủng')
p('Bác sĩ chọn thú cưng → xem lịch sử → xác định hồ sơ nhân sự → nhập chẩn đoán, điều trị hoặc thông tin vaccine → lưu bản ghi → tải lại lịch sử. Khách hàng mở sổ sức khỏe để xem bệnh án, mũi tiêm và ngày nhắc lại.')
sub('Quy trình khôi phục mật khẩu')
p('Khách nhập số điện thoại → hệ thống chuẩn hóa và tra cứu tài khoản đủ điều kiện → gửi OTP đến email liên kết → khách nhập OTP → nhận vé đặt lại có thời hạn → nhập mật khẩu mới → máy chủ đổi mật khẩu Firebase và tiêu thụ vé đặt lại.')

page(); h('F Đánh giá tiến độ và kế hoạch tiếp theo')
sub('Kết quả kiểm tra hiện tại')
p('Ngày 21/09/2026, bản dựng production của giao diện React hoàn thành thành công với 46 module được xử lý. Kết quả này xác nhận khả năng đóng gói giao diện, chưa xác nhận API, dữ liệu SQL và các tích hợp hoạt động đầy đủ trong một phiên sử dụng thực tế.')
table(['Hạng mục','Kết quả hiện tại','Điều kiện hoàn tất'],[
('Giao diện web','Đã xây dựng bốn cổng vai trò; dựng bản production đạt.','Kiểm thử thao tác trên máy tính và điện thoại.'),
('Nghiệp vụ và API','Có mã cho các luồng chính; còn điểm cần hoàn thiện.','Đạt kiểm thử phân quyền, dữ liệu và trạng thái.'),
('Tích hợp bên ngoài','Có mã Firebase, PayOS, Cloudinary và SMTP.','Chạy thành công kịch bản đầu cuối với cấu hình thực tế.'),
('Nghiệm thu và triển khai','Chưa kết luận đạt nghiệm thu.','Có biên bản kiểm thử, bản triển khai và minh chứng.')],[3.5,6.3,6.9])
sub('Kế hoạch đề xuất theo thứ tự ưu tiên')
p('Ưu tiên 1. Hoàn thiện xác thực và phân quyền tại API; kiểm tra chủ sở hữu thú cưng, lịch hẹn và hồ sơ sức khỏe. Kiểm thử bằng ít nhất hai tài khoản khách hàng và các vai trò nội bộ.')
p('Ưu tiên 2. Bảo đảm tạo lịch, chi tiết và thanh toán nhất quán; tính giá tại máy chủ; sửa cơ chế sinh mã; bổ sung kiểm tra ngày giờ và trùng lịch.')
p('Ưu tiên 3. Kiểm thử tiền mặt và PayOS, gồm webhook gửi lại, giao dịch thất bại, hủy lịch sau thanh toán và đối soát hoàn tiền. Thống nhất giới hạn ảnh và kiểm thử OTP hết hạn, sai mã, gửi lại.')
p('Ưu tiên 4. Kiểm thử toàn bộ bốn vai trò, xử lý mạng lỗi và màn hình nhỏ; hoàn thiện cấu hình triển khai, tài liệu hướng dẫn và phân công thành viên.')
sub('Nguồn đối chiếu trong dự án')
p('README.md; web-react/package.json; App.jsx, auth.jsx, api.js, CustomerPortal.jsx và OperationsPortal.jsx trong web-react/src; Program.cs, các controller nghiệp vụ, PasswordResetService.cs và ImageUploadValidator.cs trong backend/PetNoVaApi.')
p('Mốc tiếp theo: nghiệm thu một kịch bản xuyên suốt từ khách đăng nhập, thêm thú cưng, đặt lịch và trả tiền đến nhân viên hoàn tất dịch vụ, bác sĩ lưu hồ sơ và khách xem lại kết quả.')

footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=footer.add_run('PetNoVa  |  TD01  |  '); r.font.size=Pt(10)
fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); footer._p.append(fld)
doc.core_properties.title='Báo cáo tiến độ website PetNoVa'
doc.core_properties.subject='TD01'
doc.core_properties.author=''
out=root/'docs'/'TD01_NhomXX.docx'; doc.save(out)
print(out)
print('Paragraphs:',len(doc.paragraphs),'Tables:',len(doc.tables))
